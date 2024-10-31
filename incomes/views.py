import io
from django.http import HttpResponse
from django.urls import reverse_lazy, reverse
from rest_framework import permissions
from main.views import BaseDetailView, BaseCreateView, BaseUpdateView, BaseDeleteView, \
    BaseOperationViewSet, BaseOperationListView
from .forms import IncomeForm
from .models import Income
from .serializers import IncomeSerializer
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference


class IncomeViewSet(BaseOperationViewSet):
    queryset = Income.objects.all()
    serializer_class = IncomeSerializer
    permission_classes = [permissions.IsAuthenticated]


# CRUD
class IncomeListView(BaseOperationListView):
    model = Income
    template_name = 'main/operation_list.html'

    detail_url_name = 'incomes:incomes_detail'
    update_url_name = 'incomes:incomes_update'
    delete_url_name = 'incomes:incomes_delete'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['operation_type'] = 'Доход'
        return context

    def get_list_url(self):
        return reverse('incomes:incomes_list')

    def get_create_url(self):
        return reverse('incomes:incomes_create')

class IncomeDetailView(BaseDetailView):
    model = Income
    template_name = 'main/operation_detail.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['update_url'] = reverse('incomes:incomes_update', kwargs={'pk':self.object.pk})
        context['delete_url'] = reverse('incomes:incomes_delete', kwargs={'pk':self.object.pk})
        context['operation_list_url'] = reverse('incomes:incomes_list')
        return context

class IncomeCreateView(BaseCreateView):
    model = Income
    template_name = 'main/operation_form.html'
    success_url = reverse_lazy('incomes:incomes_list')
    form_class = IncomeForm

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['operation_type'] = "Доход"
        return context


class IncomeUpdateView(BaseUpdateView):
    model = Income
    template_name = 'main/operation_form.html'
    success_url = reverse_lazy('incomes:incomes_list')
    form_class = IncomeForm

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['operation_type'] = "Доход"
        return context


class IncomeDeleteView(BaseDeleteView):
    model = Income
    template_name = 'main/operation_confirm_delete.html'
    success_url = reverse_lazy('incomes:incomes_list')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['detail_url'] = reverse('incomes:incomes_detail', kwargs={'pk': self.object.pk})
        return context


# Экспорт отчетов
def export_incomes_csv(request):
    # Получаю все доходы для текущего пользователя
    incomes = Income.objects.filter(user=request.user).values()
    df = pd.DataFrame(incomes)
    df.drop(columns=['id', 'user_id'], inplace=True)
    df.rename(columns={
        'amount': 'Сумма',
        'date': 'Дата',
        'source': 'Источник',
        'category': 'Категория',
        'context': 'Комментарий'
    }, inplace=True)

    # Подсчет общей суммы доходов
    total_income = df['Сумма'].sum()
    total_row = pd.DataFrame({'Сумма': [total_income]})
    df = pd.concat([df, total_row], ignore_index=True)

    buffer = io.StringIO()
    df.to_csv(buffer, index=False)

    response = HttpResponse(buffer.getvalue(), content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="incomes.csv"'
    return response

def export_incomes_excel(request):
    incomes = Income.objects.filter(user=request.user).values()
    wb = Workbook()

    # Основной лист с данными
    ws_data = wb.active
    ws_data.title = "Доходы"
    ws_data.append(["Сумма", "Дата", "Источник", "Категория", "Комментарий"])

    for income in incomes:
        ws_data.append([income['amount'], income['date'], income['source'], income['category'], income['context']])

    # Лист с аналитикой
    ws_summary = wb.create_sheet(title='Аналитика')
    total_income = sum([income['amount'] for income in incomes])
    ws_summary.append(['Общая сумма доходов:', total_income])

    # Диаграмма по категориям
    categories = {}
    for income in incomes:
        categories[income['category']] = categories.get(income['category'], 0) + income['amount']

    ws_summary.append(['Категория', 'Сумма'])
    for category, amount in categories.items():
        ws_summary.append([category, amount])

    pie = PieChart()
    labels = Reference(ws_summary, min_col=1, min_row=2, max_row=1 + len(categories))
    data = Reference(ws_summary, min_col=2, min_row=2, max_row=1 + len(categories))
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    pie.title = 'Доходы по категориям'
    ws_summary.add_chart(pie, 'D5')  # Позиция диаграммы

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="incomes.xlsx"'
    wb.save(response)
    return response